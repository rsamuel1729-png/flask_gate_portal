import os, sqlite3, time, random, urllib.request, xml.etree.ElementTree as ET, csv, io, re, json, uuid
from datetime import date, datetime, timedelta
from functools import wraps
from flask import Flask, g, render_template, request, redirect, url_for, session, flash, abort, jsonify
from werkzeug.security import generate_password_hash, check_password_hash
from pymongo import MongoClient, ReturnDocument
from pymongo.errors import PyMongoError

app = Flask(__name__)
app.config.update(
    SECRET_KEY=os.environ.get("SECRET_KEY", "change-me-in-production"),
    MONGO_URI=os.environ.get("MONGO_URI", "mongodb://127.0.0.1:27017"),
    MONGO_DATABASE=os.environ.get("MONGO_DATABASE", "gate_portal"),
    LEGACY_DATABASE=os.path.join(app.root_path, "gate.db"),
    GATE_EXAM_DATE=os.environ.get("GATE_EXAM_DATE", "2027-02-06"),  # verify on the official GATE site
    GATE_FEED_URL=os.environ.get("GATE_FEED_URL", ""),               # optional RSS/Atom feed for live news
    ADMIN_EMAIL=os.environ.get("ADMIN_EMAIL", "admin@gate.local"),
    ADMIN_PASSWORD=os.environ.get("ADMIN_PASSWORD", "admin123"),
)

# ---------- MongoDB ----------
mongo_client = MongoClient(app.config["MONGO_URI"], serverSelectionTimeoutMS=5000)
mongo = mongo_client[app.config["MONGO_DATABASE"]]

def collection(name): return mongo[name]

def public_doc(doc):
    if not doc: return None
    doc = dict(doc)
    doc.setdefault("id", doc.get("_id"))
    doc.pop("_id", None)
    return doc

def find_many(name, filter=None, sort=None, limit=None):
    cur = collection(name).find(filter or {})
    if sort: cur = cur.sort(sort)
    if limit: cur = cur.limit(limit)
    return [public_doc(doc) for doc in cur]

def find_one(name, filter=None): return public_doc(collection(name).find_one(filter or {}))

def add_doc(name, data):
    data = dict(data)
    if "id" not in data:
        counter = mongo["_counters"].find_one_and_update({"_id": name}, {"$inc": {"seq": 1}}, upsert=True, return_document=ReturnDocument.AFTER)
        data["id"] = counter["seq"]
    data["_id"] = data["id"]
    collection(name).insert_one(data)
    return data["id"]

def put_doc(name, filter, data): collection(name).update_one(filter, {"$set": data}, upsert=True)
def remove_doc(name, filter): return collection(name).delete_one(filter).deleted_count
def count_docs(name, filter=None): return collection(name).count_documents(filter or {})

def yt(s): return "https://www.youtube.com/results?search_query=" + s.replace(" ", "+")

def read_question_document(filename, data):
    ext = os.path.splitext(filename.lower())[1]
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif ext == ".docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        parts.extend("\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows) for table in doc.tables)
        text = "\n".join(parts)
    else:
        raise ValueError("Use a text-based PDF or .docx Word document. Legacy .doc files are not supported.")
    if not text.strip():
        raise ValueError("No selectable text found. This may be a scanned PDF; OCR it or save it as a text-based PDF first.")
    return text

def parse_question_document(text):
    start_re = re.compile(r"^\s*(?:Q(?:uestion)?\s*)?(\d{1,4})\s*[.)]\s*(.+)$", re.I | re.M)
    starts = list(start_re.finditer(text))
    parsed, issues = [], []
    for index, match in enumerate(starts):
        block = text[match.end(): starts[index + 1].start() if index + 1 < len(starts) else len(text)]
        block = match.group(2) + "\n" + block
        option_re = re.compile(r"^\s*([A-D])\s*[).:]\s*(.*)$", re.I | re.M)
        opts = list(option_re.finditer(block))
        if len(opts) < 4:
            if opts or len(match.group(2).strip()) > 10:
                issues.append(f"Question {match.group(1)} has fewer than four recognizable options; it was not imported.")
            continue  # Numbered answer-key lines are not questions.
        values, preamble = {}, block[:opts[0].start()]
        for oi, om in enumerate(opts):
            end = opts[oi + 1].start() if oi + 1 < len(opts) else len(block)
            value = (om.group(2) + " " + block[om.end():end]).strip()
            value = re.split(r"\n\s*(?:answer|ans|correct answer|explanation|solution|topic)\s*[:=-]", value, maxsplit=1, flags=re.I)[0]
            values[om.group(1).upper()] = re.sub(r"\s+", " ", value).strip()
        answer = re.search(r"(?:correct\s+answer|answer|ans)\s*(?:is|[:=-])?\s*[\[(]?([A-D])\b", block, re.I)
        expl = re.search(r"(?:explanation|solution)\s*[:=-]\s*(.*?)(?=\n\s*(?:topic|Q(?:uestion)?\s*\d+\s*[.)])|$)", block, re.I | re.S)
        topic = re.search(r"topic\s*[:=-]\s*(.+)", block, re.I)
        question = re.sub(r"\s+", " ", preamble).strip()
        if not question or set(values) != {"A", "B", "C", "D"} or not answer:
            issues.append(f"Question {match.group(1)} was found but needs question text, A-D options, and an explicit Answer: A/B/C/D.")
            continue
        parsed.append({"number": int(match.group(1)), "text": question, "a": values["A"], "b": values["B"], "c": values["C"], "d": values["D"],
                       "ans": answer.group(1).upper(), "expl": re.sub(r"\s+", " ", expl.group(1)).strip() if expl else "",
                       "topic": re.sub(r"\s+", " ", topic.group(1)).strip() if topic else ""})
    if not starts:
        issues.append("No numbered questions found. Use headings like Q1. Question text, then A) through D) options and Answer: B.")
    elif not parsed and not issues:
        issues.append("No complete multiple-choice questions found. Check the document structure and include the correct answer for each question.")
    return parsed, issues

SEED = {  # subject: (weightage, icon, {topic: [(title,url,kind)]})
 "Engineering Mathematics": (13, "∑", {
   "Discrete Mathematics: Logic & Sets": [("NPTEL Discrete Mathematics", "https://nptel.ac.in/courses/106106183", "video")],
   "Graph Theory & Combinatorics": [("Graph Theory lectures", yt("graph theory gate cse"), "video")],
   "Linear Algebra": [("Gilbert Strang MIT 18.06", "https://ocw.mit.edu/courses/18-06-linear-algebra-spring-2010/", "notes")],
   "Probability & Statistics": [("Probability for GATE", yt("probability gate cse"), "video")],
   "Calculus": [("MIT Single Variable Calculus", "https://ocw.mit.edu/courses/18-01sc-single-variable-calculus-fall-2010/", "notes")]}),
 "Digital Logic": (6, "⚙", {
   "Boolean Algebra & K-maps": [("Digital Logic basics", yt("digital logic boolean algebra kmap gate"), "video")],
   "Combinational Circuits": [("Combinational circuits", yt("combinational circuits gate cse"), "video")],
   "Sequential Circuits": [("Flip-flops & counters", yt("sequential circuits flip flops gate cse"), "video")],
   "Number Systems": [("Number systems", yt("number systems floating point gate"), "video")]}),
 "Computer Organization & Architecture": (8, "🖥", {
   "Machine Instructions & Addressing": [("COA lectures", yt("computer organization addressing modes gate"), "video")],
   "Pipelining": [("Pipelining hazards", yt("pipelining hazards gate cse"), "video")],
   "Cache & Memory Hierarchy": [("Cache memory", yt("cache memory mapping gate cse"), "video")],
   "I/O Interface & DMA": [("I/O and DMA", yt("DMA interrupts gate cse"), "video")]}),
 "Programming & Data Structures": (10, "{ }", {
   "C Programming & Recursion": [("C programming NPTEL", "https://nptel.ac.in/courses/106105171", "video")],
   "Arrays, Stacks, Queues, Linked Lists": [("Data structures basics", yt("stack queue linked list gate cse"), "video")],
   "Trees, BST, Heaps": [("Trees and heaps", yt("binary tree heap gate cse"), "video")],
   "Hashing & Graphs": [("Hashing", yt("hashing gate cse"), "video")]}),
 "Algorithms": (9, "⚡", {
   "Asymptotic Analysis": [("Asymptotic notation", yt("asymptotic notations gate cse"), "video")],
   "Divide & Conquer, Sorting": [("MIT 6.006 Intro to Algorithms", "https://ocw.mit.edu/courses/6-006-introduction-to-algorithms-spring-2020/", "notes")],
   "Greedy & Dynamic Programming": [("DP for GATE", yt("dynamic programming gate cse"), "video")],
   "Graph Algorithms (BFS, DFS, MST, SSSP)": [("Graph algorithms", yt("graph algorithms mst dijkstra gate cse"), "video")]}),
 "Theory of Computation": (9, "∞", {
   "Regular Languages & Finite Automata": [("TOC NPTEL", "https://nptel.ac.in/courses/106104028", "video")],
   "Context-Free Languages & PDA": [("CFG and PDA", yt("context free grammar pda gate cse"), "video")],
   "Turing Machines & Decidability": [("Decidability", yt("turing machine decidability gate cse"), "video")]}),
 "Compiler Design": (5, "⌘", {
   "Lexical Analysis & Parsing": [("Parsing techniques", yt("LL1 LR parsing gate cse"), "video")],
   "Syntax Directed Translation": [("SDT", yt("syntax directed translation gate cse"), "video")],
   "Code Generation & Optimization": [("Code optimization", yt("code optimization compiler gate cse"), "video")]}),
 "Operating Systems": (10, "🧠", {
   "Processes & Threads": [("OS lectures", yt("process thread operating system gate cse"), "video")],
   "CPU Scheduling": [("Scheduling algorithms", yt("cpu scheduling gate cse"), "video")],
   "Synchronization & Deadlocks": [("Semaphores & deadlock", yt("semaphore deadlock gate cse"), "video")],
   "Memory Management & Virtual Memory": [("Paging and page replacement", yt("paging virtual memory gate cse"), "video")],
   "File Systems & Disk": [("File systems", yt("file system disk scheduling gate cse"), "video")]}),
 "Databases": (9, "🗄", {
   "ER Model & Relational Algebra": [("DBMS NPTEL", "https://nptel.ac.in/courses/106105175", "video")],
   "SQL": [("SQL practice", yt("sql queries gate cse"), "video")],
   "Normalization & Functional Dependencies": [("Normalization", yt("normalization functional dependency gate cse"), "video")],
   "Transactions & Concurrency Control": [("Transactions", yt("transaction serializability gate cse"), "video")],
   "Indexing & B/B+ Trees": [("B+ trees", yt("b+ tree indexing gate cse"), "video")]}),
 "Computer Networks": (8, "🌐", {
   "OSI/TCP-IP & Data Link": [("CN NPTEL", "https://nptel.ac.in/courses/106105183", "video")],
   "IP Addressing & Routing": [("Subnetting & routing", yt("subnetting routing gate cse"), "video")],
   "Transport Layer: TCP/UDP": [("TCP congestion control", yt("tcp congestion control gate cse"), "video")],
   "Application Layer Protocols": [("DNS, HTTP, SMTP", yt("dns http smtp gate cse"), "video")]}),
 "General Aptitude": (15, "✎", {
   "Verbal Ability": [("Verbal for GATE", yt("gate verbal ability"), "video")],
   "Quantitative Aptitude": [("Quant for GATE", yt("gate quantitative aptitude"), "video")],
   "Analytical & Spatial Reasoning": [("Reasoning", yt("gate analytical spatial reasoning"), "video")]}),
}
SEED_Q = {  # subject: [(question, a, b, c, d, ans, explanation)]
 "Operating Systems": [
  ("Which scheduling algorithm can cause starvation of long processes?", "FCFS", "Round Robin", "Shortest Job First", "None", "C", "SJF keeps favouring short jobs, so long ones may wait indefinitely."),
  ("Which condition is NOT necessary for deadlock?", "Mutual exclusion", "Hold and wait", "Preemption", "Circular wait", "C", "No preemption is required; preemption would break deadlock."),
  ("With 4 frames and FIFO, Belady's anomaly is seen when:", "Frames decrease faults", "Frames increase faults", "LRU is used", "Pages are sorted", "B", "FIFO can show more faults with more frames.")],
 "Databases": [
  ("A relation in 3NF is always in:", "BCNF", "2NF", "4NF", "5NF", "B", "Normal forms are nested: 3NF implies 2NF."),
  ("Which SQL clause filters groups?", "WHERE", "HAVING", "GROUP BY", "ORDER BY", "B", "HAVING applies after aggregation."),
  ("B+ tree leaf nodes are:", "Not linked", "Linked sequentially", "Always root", "Hashed", "B", "Leaves are linked for range queries.")],
 "Algorithms": [
  ("Time complexity of Merge Sort (worst case)?", "O(n)", "O(n log n)", "O(n^2)", "O(log n)", "B", "T(n)=2T(n/2)+n gives n log n."),
  ("Dijkstra's algorithm fails with:", "Cycles", "Negative edge weights", "Dense graphs", "Directed graphs", "B", "Negative edges break the greedy invariant."),
  ("0/1 Knapsack is best solved using:", "Greedy", "Dynamic programming", "BFS", "Divide only", "B", "Greedy by ratio fails for 0/1 version.")],
 "Computer Networks": [
  ("Number of usable hosts in a /26 network?", "62", "64", "30", "126", "A", "2^6 - 2 = 62."),
  ("TCP uses which mechanism for flow control?", "Sliding window", "Token bucket", "CSMA/CD", "ARP", "A", "Receiver window advertises buffer space."),
  ("Which layer is responsible for routing?", "Data link", "Network", "Transport", "Session", "B", "The network layer routes packets.")],
 "Digital Logic": [
  ("Minimum number of NAND gates to implement XOR?", "3", "4", "5", "2", "B", "XOR needs 4 two-input NAND gates."),
  ("A 3-to-8 decoder has how many outputs?", "3", "6", "8", "16", "C", "2^3 = 8 outputs.")],
 "Theory of Computation": [
  ("Which is NOT closed under complement?", "Regular", "Deterministic CFL", "Context-free", "Recursive", "C", "CFLs are not closed under complement."),
  ("The halting problem is:", "Decidable", "Undecidable", "Regular", "Context-free", "B", "Turing proved it undecidable.")],
 "Engineering Mathematics": [
  ("Number of edges in a complete graph K_n?", "n", "n(n-1)/2", "n^2", "2^n", "B", "Choose any 2 vertices."),
  ("P(A∪B) for mutually exclusive events equals:", "P(A)P(B)", "P(A)+P(B)", "P(A)-P(B)", "1", "B", "Addition rule with zero intersection.")],
 "Programming & Data Structures": [
  ("Inorder traversal of a BST gives:", "Random order", "Sorted order", "Reverse order", "Level order", "B", "Left-root-right visits ascending."),
  ("Stack is a ___ structure.", "FIFO", "LIFO", "Random", "Hashed", "B", "Last in, first out.")],
 "Computer Organization & Architecture": [
  ("Which hazard arises from dependent instructions?", "Structural", "Data", "Control", "None", "B", "Data hazards: RAW, WAR, WAW."),
  ("Direct-mapped cache with 8 lines: block 12 maps to line:", "2", "4", "5", "0", "B", "12 mod 8 = 4.")],
}

# ---------- MongoDB startup and one-time SQLite migration ----------
def init_db():
    mongo_client.admin.command("ping")
    collection("users").create_index("email", unique=True)
    collection("progress").create_index([("user_id", 1), ("topic_id", 1)], unique=True)
    collection("bookmarks").create_index([("user_id", 1), ("resource_id", 1)], unique=True)
    collection("notes").create_index([("user_id", 1), ("topic_id", 1)], unique=True)
    collection("attempt_answers").create_index([("attempt_id", 1), ("question_id", 1)], unique=True)
    if count_docs("subjects") == 0 and os.path.exists(app.config["LEGACY_DATABASE"]):
        legacy = sqlite3.connect(app.config["LEGACY_DATABASE"])
        legacy.row_factory = sqlite3.Row
        tables = [r[0] for r in legacy.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'")]
        for table in tables:
            docs = []
            for i, row in enumerate(legacy.execute(f'SELECT * FROM "{table}"')):
                doc = dict(row)
                if "id" in doc and doc["id"] is not None:
                    doc["_id"] = doc["id"]
                else:
                    doc["_id"] = f"legacy:{table}:{i}"
                docs.append(doc)
            if docs:
                collection(table).insert_many(docs, ordered=False)
                numeric = [d.get("id") for d in docs if isinstance(d.get("id"), int)]
                if numeric:
                    mongo["_counters"].update_one({"_id": table}, {"$max": {"seq": max(numeric)}}, upsert=True)
        legacy.close()
    if count_docs("subjects") == 0:
        topic_ids = {}
        for name, (weight, icon, topic_map) in SEED.items():
            sid = add_doc("subjects", {"name": name, "weightage": weight, "icon": icon})
            topic_ids[name] = {}
            for topic_name, resources in topic_map.items():
                tid = add_doc("topics", {"subject_id": sid, "name": topic_name})
                topic_ids[name][topic_name] = tid
                for title, link, kind in resources:
                    add_doc("resources", {"topic_id": tid, "title": title, "url": link, "kind": kind})
        rules = {
            "Operating Systems": [("scheduling", "CPU Scheduling"), ("deadlock", "Synchronization & Deadlocks"), ("frames", "Memory Management & Virtual Memory")],
            "Databases": [("3nf", "Normalization & Functional Dependencies"), ("sql clause", "SQL"), ("b+ tree", "Indexing & B/B+ Trees")],
            "Algorithms": [("merge sort", "Divide & Conquer, Sorting"), ("dijkstra", "Graph Algorithms (BFS, DFS, MST, SSSP)"), ("knapsack", "Greedy & Dynamic Programming")],
            "Computer Networks": [("/26", "IP Addressing & Routing"), ("tcp", "Transport Layer: TCP/UDP"), ("routing", "IP Addressing & Routing")],
            "Digital Logic": [("nand", "Boolean Algebra & K-maps"), ("decoder", "Combinational Circuits")],
            "Theory of Computation": [("closed", "Context-Free Languages & PDA"), ("halting", "Turing Machines & Decidability")],
            "Engineering Mathematics": [("complete graph", "Graph Theory & Combinatorics"), ("mutually exclusive", "Probability & Statistics")],
            "Programming & Data Structures": [("bst", "Trees, BST, Heaps"), ("stack", "Arrays, Stacks, Queues, Linked Lists")],
            "Computer Organization & Architecture": [("hazard", "Pipelining"), ("cache", "Cache & Memory Hierarchy")],
        }
        for name, questions in SEED_Q.items():
            sid = find_one("subjects", {"name": name})["id"]
            for question in questions:
                text = question[0]
                topic_id = next((topic_ids[name][topic] for needle, topic in rules.get(name, []) if needle in text.lower()), None)
                add_doc("questions", dict(zip(("text", "a", "b", "c", "d", "ans", "expl"), question), subject_id=sid, topic_id=topic_id,
                                               source="Original practice", year=None, kind="MCQ", important=0, bank_id=None, bank_set=None))
    if count_docs("updates") == 0:
        add_doc("updates", {"title": "Welcome to GATE CSE Prep", "body": "Admins can post exam notifications here. Always confirm dates on the official GATE website.", "link": "https://gate2026.iitg.ac.in", "created": datetime.now().isoformat()})
    if count_docs("users", {"is_admin": 1}) == 0:
        add_doc("users", {"name": "Admin", "email": app.config["ADMIN_EMAIL"], "pw": generate_password_hash(app.config["ADMIN_PASSWORD"]), "is_admin": 1, "hours": 4, "exam_date": app.config["GATE_EXAM_DATE"]})

# ---------- Auth ----------
def login_required(f):
    @wraps(f)
    def w(*a, **k):
        if "uid" not in session: return redirect(url_for("login", next=request.path))
        return f(*a, **k)
    return w

def admin_required(f):
    @wraps(f)
    @login_required
    def w(*a, **k):
        if not g.user["is_admin"]: abort(403)
        return f(*a, **k)
    return w

@app.before_request
def load_user():
    g.user = find_one("users", {"id": session["uid"]}) if "uid" in session else None

@app.context_processor
def inject():
    d = (datetime.strptime(g.user["exam_date"] or app.config["GATE_EXAM_DATE"], "%Y-%m-%d").date() - date.today()).days if g.user else None
    return dict(days_left=d)

@app.route("/register", methods=["GET", "POST"])
def register():
    if request.method == "POST":
        n, e, p = request.form["name"].strip(), request.form["email"].strip().lower(), request.form["password"]
        if not n or "@" not in e or len(p) < 6:
            flash("Enter a name, valid email and a password of 6+ characters.", "err")
        elif find_one("users", {"email": e}):
            flash("Email already registered.", "err")
        else:
            session["uid"] = add_doc("users", {"name": n, "email": e, "pw": generate_password_hash(p), "exam_date": app.config["GATE_EXAM_DATE"], "is_admin": 0, "hours": 4})
            return redirect(url_for("dashboard"))
    return render_template("auth.html", mode="register")

@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        u = find_one("users", {"email": request.form["email"].strip().lower()})
        if u and check_password_hash(u["pw"], request.form["password"]):
            session["uid"] = u["id"]
            nxt = request.args.get("next", "")
            return redirect(nxt if nxt.startswith("/") and not nxt.startswith("//") else url_for("dashboard"))
        flash("Invalid email or password.", "err")
    return render_template("auth.html", mode="login")

@app.route("/admin/login", methods=["GET", "POST"])
def admin_login():
    if g.user and g.user["is_admin"]:
        return redirect(url_for("admin"))
    if request.method == "POST":
        u = find_one("users", {"email": request.form["email"].strip().lower()})
        if u and u["is_admin"] and check_password_hash(u["pw"], request.form["password"]):
            session["uid"] = u["id"]
            return redirect(url_for("admin"))
        flash("Admin sign-in failed. Check the admin email and password.", "err")
    return render_template("auth.html", mode="admin")

@app.route("/logout")
def logout():
    session.clear(); return redirect(url_for("index"))

# ---------- Pages ----------
@app.route("/")
def index():
    if g.user: return redirect(url_for("dashboard"))
    return render_template("index.html", updates=find_many("updates", sort=[("id", -1)], limit=3))

def subject_stats(uid):
    result = []
    for subject in find_many("subjects", sort=[("id", 1)]):
        topics = find_many("topics", {"subject_id": subject["id"]})
        completed = {p["topic_id"] for p in find_many("progress", {"user_id": uid})}
        subject.update(total=len(topics), done=sum(t["id"] in completed for t in topics))
        result.append(subject)
    return result

@app.route("/dashboard")
@login_required
def dashboard():
    subs = subject_stats(g.user["id"])
    total, done = sum(s["total"] for s in subs), sum(s["done"] for s in subs)
    attempts = find_many("attempts", {"user_id": g.user["id"]}, sort=[("id", -1)], limit=5)
    for attempt in attempts:
        subject = find_one("subjects", {"id": attempt.get("subject_id")}) if attempt.get("subject_id") else None
        attempt["sname"] = subject["name"] if subject else None
    return render_template("dashboard.html", subs=subs, total=total, done=done, attempts=attempts,
                           updates=find_many("updates", sort=[("id", -1)], limit=3),
                           bookmarks=[r for b in find_many("bookmarks", {"user_id": g.user["id"]}, limit=5) if (r := find_one("resources", {"id": b["resource_id"]}))])

@app.route("/subjects")
@login_required
def subjects():
    return render_template("subjects.html", subs=subject_stats(g.user["id"]))

@app.route("/subject/<int:sid>")
@login_required
def subject(sid):
    s = find_one("subjects", {"id": sid}) or abort(404)
    uid = g.user["id"]
    topics = find_many("topics", {"subject_id": sid}, sort=[("id", 1)])
    progress_ids = {p["topic_id"] for p in find_many("progress", {"user_id": uid})}
    for topic in topics:
        topic["done"] = topic["id"] in progress_ids
        note = find_one("notes", {"user_id": uid, "topic_id": topic["id"]})
        topic["note"] = note.get("body") if note else None
    bookmarked = {b["resource_id"] for b in find_many("bookmarks", {"user_id": uid})}
    res = {}
    for topic in topics:
        res[topic["id"]] = find_many("resources", {"topic_id": topic["id"]})
        for resource in res[topic["id"]]: resource["bm"] = resource["id"] in bookmarked
    return render_template("subject.html", s=s, topics=topics, res=res)

@app.post("/topic/<int:tid>/toggle")
@login_required
def toggle(tid):
    t = find_one("topics", {"id": tid}) or abort(404)
    if find_one("progress", {"user_id": g.user["id"], "topic_id": tid}):
        remove_doc("progress", {"user_id": g.user["id"], "topic_id": tid})
    else:
        put_doc("progress", {"user_id": g.user["id"], "topic_id": tid}, {"user_id": g.user["id"], "topic_id": tid, "done_at": datetime.now().isoformat()})
    return redirect(url_for("subject", sid=t["subject_id"]) + f"#t{tid}")

@app.post("/topic/<int:tid>/note")
@login_required
def note(tid):
    t = find_one("topics", {"id": tid}) or abort(404)
    put_doc("notes", {"user_id": g.user["id"], "topic_id": tid}, {"user_id": g.user["id"], "topic_id": tid, "body": request.form.get("body", "")[:2000]})
    flash("Note saved.", "ok")
    return redirect(url_for("subject", sid=t["subject_id"]) + f"#t{tid}")

@app.post("/resource/<int:rid>/bookmark")
@login_required
def bookmark(rid):
    r = find_one("resources", {"id": rid}) or abort(404)
    topic = find_one("topics", {"id": r["topic_id"]}) or abort(404)
    if find_one("bookmarks", {"user_id": g.user["id"], "resource_id": rid}):
        remove_doc("bookmarks", {"user_id": g.user["id"], "resource_id": rid})
    else:
        put_doc("bookmarks", {"user_id": g.user["id"], "resource_id": rid}, {"user_id": g.user["id"], "resource_id": rid})
    return redirect(url_for("subject", sid=topic["subject_id"]))

@app.route("/search")
@login_required
def search():
    term = request.args.get("q", "").strip()
    pattern = re.compile(re.escape(term), re.I)
    topics = find_many("topics") if term else []
    topics = [dict(t, sname=(find_one("subjects", {"id": t["subject_id"]}) or {}).get("name", "")) for t in topics if pattern.search(t.get("name", ""))]
    resources = find_many("resources") if term else []
    res = []
    for resource in resources:
        if pattern.search(resource.get("title", "")):
            topic = find_one("topics", {"id": resource["topic_id"]})
            res.append(dict(resource, subject_id=topic["subject_id"] if topic else None))
    return render_template("search.html", term=term, topics=topics, res=res)

# ---------- Study plan ----------
@app.route("/plan", methods=["GET", "POST"])
@login_required
def plan():
    if request.method == "POST":
        try:
            ed = datetime.strptime(request.form["exam_date"], "%Y-%m-%d").date(); h = float(request.form["hours"])
            assert 0.5 <= h <= 16
        except Exception:
            flash("Enter a valid exam date and 0.5-16 hours/day.", "err"); return redirect(url_for("plan"))
        collection("users").update_one({"id": g.user["id"]}, {"$set": {"exam_date": ed.isoformat(), "hours": h}})
        return redirect(url_for("plan"))
    ed = datetime.strptime(g.user["exam_date"] or app.config["GATE_EXAM_DATE"], "%Y-%m-%d").date()
    days = (ed - date.today()).days
    completed = {p["topic_id"] for p in find_many("progress", {"user_id": g.user["id"]})}
    rows = []
    for topic in find_many("topics"):
        if topic["id"] in completed: continue
        subject = find_one("subjects", {"id": topic["subject_id"]})
        if subject: rows.append(dict(topic, sname=subject["name"], weightage=subject["weightage"]))
    rows.sort(key=lambda item: (-item["weightage"], item["id"]))
    hrs, per_topic = g.user["hours"], 4.0       # estimated study and practice effort per topic
    revision_days = max(0, min(30, days // 4))  # last 25% (max 30 days) reserved for revision & mocks
    study_days = max(0, days - revision_days)
    cap = study_days * hrs
    need = len(rows) * per_topic
    # Split topic effort across days when a student has fewer than four hours/day.
    schedule, d = [], date.today()
    day_items, left = [], 0.0
    for t in rows:
        remaining_effort = per_topic
        while remaining_effort > 0 and d <= ed and d < date.today() + timedelta(days=study_days):
            chunk = min(remaining_effort, hrs - left)
            day_items.append({"topic": t, "hours": round(chunk, 1)})
            left = round(left + chunk, 1)
            remaining_effort = round(remaining_effort - chunk, 1)
            if left >= hrs:
                schedule.append((d, day_items))
                day_items, left, d = [], 0.0, d + timedelta(days=1)
    if day_items: schedule.append((d, day_items))
    return render_template("plan.html", days=days, ed=ed, hrs=hrs, need=need, cap=cap, revision_days=revision_days,
                           schedule=schedule[:21], remaining=len(rows), feasible=need <= cap,
                           suggest=round(need / study_days, 1) if study_days else None)

# ---------- Mock tests ----------
def choose_mock_questions(uid, subject_id=None, limit=30):
    attempts = find_many("attempts", {"user_id": uid})
    attempt_ids = {a["id"] for a in attempts}
    used_answers = find_many("attempt_answers", {"attempt_id": {"$in": list(attempt_ids)}}) if attempt_ids else []
    seen_ids = {a["question_id"] for a in used_answers}
    topics = {t["id"]: t["name"] for t in find_many("topics")}
    banks = {b["id"]: b for b in find_many("question_banks")}
    pool = find_many("questions", {"subject_id": subject_id}) if subject_id is not None else find_many("questions")
    for item in pool:
        item["topic_name"] = topics.get(item.get("topic_id"))
        bank = banks.get(item.get("bank_id"))
        item["bank_name"] = bank.get("name") if bank else None
        item["bank_created"] = bank.get("created", "9999") if bank else "9999"
    pool.sort(key=lambda item: (item["subject_id"], item["bank_created"], item.get("bank_set") or 999999, item["id"]))
    unseen = [item for item in pool if item["id"] not in seen_ids]
    def mixed_order(items):
        if subject_id is not None:
            return items
        groups = {}
        for row in items:
            groups.setdefault(row["subject_id"], []).append(row)
        output = []
        while groups:
            for key in list(groups):
                output.append(groups[key].pop(0))
                if not groups[key]: del groups[key]
        return output
    unseen = mixed_order(unseen)
    if len(unseen) >= limit:
        selected = unseen[:limit]  # imported set order means 90 questions become three sequential 30s
    else:
        selected = list(unseen)
        selected_ids = {r["id"] for r in selected}
        seen = [r for r in pool if r["id"] in seen_ids and r["id"] not in selected_ids]
        if subject_id is not None:
            random.shuffle(seen)
        else:
            seen = mixed_order(seen)
        selected.extend(seen[:limit - len(selected)])
    # Never expose answer keys or explanations in the in-progress test HTML.
    selected = [dict(item) for item in selected]
    for item in selected:
        item.pop("ans", None); item.pop("expl", None)
    return selected

@app.route("/tests")
@login_required
def tests():
    subs = find_many("subjects", sort=[("id", 1)])
    for subject in subs: subject["nq"] = count_docs("questions", {"subject_id": subject["id"]})
    hist = find_many("attempts", {"user_id": g.user["id"]}, sort=[("id", -1)], limit=15)
    for attempt in hist:
        subject = find_one("subjects", {"id": attempt.get("subject_id")}) if attempt.get("subject_id") else None
        attempt["sname"] = subject["name"] if subject else None
    now = datetime.now()
    day_start = datetime.combine(now.date(), datetime.min.time()).isoformat()
    attempts = find_many("attempts", {"user_id": g.user["id"]}, sort=[("id", -1)])
    used = sum(a["taken"] >= day_start for a in attempts)
    last = attempts[0] if attempts else None
    ready_at = datetime.fromisoformat(last["taken"]) + timedelta(minutes=40) if last else now
    wait_seconds = max(0, int((ready_at - now).total_seconds())) if used < 5 else max(0, int((datetime.combine(now.date() + timedelta(days=1), datetime.min.time()) - now).total_seconds()))
    return render_template("tests.html", subs=subs, hist=hist, total=count_docs("questions"),
                           used=used, daily_limit=5, wait_seconds=wait_seconds, ready_at=ready_at)

@app.route("/test/<sid>")
@login_required
def take_test(sid):
    now = datetime.now()
    day_start = datetime.combine(now.date(), datetime.min.time()).isoformat()
    attempts = find_many("attempts", {"user_id": g.user["id"]}, sort=[("id", -1)])
    used = sum(a["taken"] >= day_start for a in attempts)
    last = attempts[0] if attempts else None
    if used >= 5:
        flash("You have reached today's five mock limit. Review your answers and come back tomorrow.", "err"); return redirect(url_for("tests"))
    if last and (now - datetime.fromisoformat(last["taken"])).total_seconds() < 2400:
        flash("Your next mock unlocks after the 40-minute recovery interval. Use Quick Revision while you wait.", "err"); return redirect(url_for("tests"))
    if sid == "full":
        qs = choose_mock_questions(g.user["id"], limit=30); title = "CSE Mixed Mock"
    else:
        try: sid_int = int(sid)
        except ValueError: abort(404)
        s = find_one("subjects", {"id": sid_int}) or abort(404)
        qs = choose_mock_questions(g.user["id"], subject_id=s["id"], limit=30); title = s["name"]
    if not qs:
        flash("No questions yet for this subject.", "err"); return redirect(url_for("tests"))
    active = session.get("active_tests", {})
    active[sid] = {"ids": [int(item["id"]) for item in qs], "started": time.time()}
    session["active_tests"] = active
    return render_template("test.html", qs=qs, title=title, sid=sid, minutes=40)

@app.post("/test/<sid>/submit")
@login_required
def submit_test(sid):
    ids = list(dict.fromkeys(int(i) for i in request.form.get("ids", "").split(",") if i.isdigit()))[:30]
    if not ids: abort(400)
    active = session.get("active_tests", {})
    issued = active.get(sid)
    if not issued or ids != issued["ids"]:
        abort(400)
    active.pop(sid, None); session["active_tests"] = active
    marks = {r["id"]: r for r in find_many("questions", {"id": {"$in": ids}})}
    c = w = 0; review = []
    for i in ids:
        qu = marks.get(i)
        if not qu: continue
        if sid == "full" or str(qu["subject_id"]) == str(sid): pass
        else: abort(400)
        a = request.form.get(f"q{i}")
        if a == qu["ans"]: c += 1
        elif a: w += 1
        topic = find_one("topics", {"id": qu["topic_id"]}) if qu.get("topic_id") else None
        review.append({"question": qu, "answer": a, "topic": topic["name"] if topic else "Uncategorized"})
    score = round(c - w / 3, 2)   # +1 correct, -1/3 wrong (GATE 1-mark MCQ style)
    aid = add_doc("attempts", {"user_id": g.user["id"], "subject_id": None if sid == "full" else int(sid), "score": score, "total": len(review),
                                "correct": c, "wrong": w, "taken": datetime.now().isoformat()})
    for item in review:
        qu, answer = item["question"], item["answer"]
        add_doc("attempt_answers", {"attempt_id": aid, "question_id": qu["id"], "selected": answer, "is_correct": int(answer == qu["ans"])})
    topic_stats = {}
    for item in review:
        label = item["topic"]
        stat = topic_stats.setdefault(label, {"total": 0, "wrong": 0, "correct": 0})
        stat["total"] += 1
        stat["correct"] += int(item["answer"] == item["question"]["ans"])
        stat["wrong"] += int(item["answer"] != item["question"]["ans"])
    weak = sorted((dict(name=k, **v, pct=round(v["correct"] * 100 / v["total"])) for k,v in topic_stats.items()), key=lambda x: x["pct"])
    user_attempts = find_many("attempts", {"user_id": g.user["id"]})
    attempt_ids = [a["id"] for a in user_attempts]
    answers = find_many("attempt_answers", {"attempt_id": {"$in": attempt_ids}}) if attempt_ids else []
    question_map = {qu["id"]: qu for qu in find_many("questions", {"id": {"$in": [a["question_id"] for a in answers]}})} if answers else {}
    topic_results = {}
    for answer in answers:
        qu = question_map.get(answer["question_id"])
        topic = find_one("topics", {"id": qu.get("topic_id")}) if qu and qu.get("topic_id") else None
        if not topic: continue
        stat = topic_results.setdefault(topic["name"], {"total": 0, "correct": 0})
        stat["total"] += 1; stat["correct"] += answer["is_correct"]
    proficiency = [dict(name=name, total=stat["total"], correct=stat["correct"], pct=round(stat["correct"]*100/stat["total"])) for name,stat in topic_results.items()]
    proficiency.sort(key=lambda x: x["pct"])
    return render_template("result.html", review=review, score=score, c=c, w=w, n=len(review), topics=weak,
                           proficiency=proficiency, attempt_id=aid)

@app.route("/revision")
@login_required
def revision():
    attempts = find_many("attempts", {"user_id": g.user["id"]})
    attempt_ids = [a["id"] for a in attempts]
    answers = find_many("attempt_answers", {"attempt_id": {"$in": attempt_ids}, "is_correct": 0}, sort=[("id", -1)], limit=60) if attempt_ids else []
    misses = {}
    rows = []
    for answer in answers:
        question = find_one("questions", {"id": answer["question_id"]})
        if not question: continue
        misses[question["id"]] = misses.get(question["id"], 0) + 1
        topic = find_one("topics", {"id": question.get("topic_id")}) if question.get("topic_id") else None
        subject = find_one("subjects", {"id": question["subject_id"]})
        attempt = find_one("attempts", {"id": answer["attempt_id"]})
        rows.append(dict(question, topic_name=topic["name"] if topic else None, subject_name=subject["name"] if subject else None,
                         selected=answer.get("selected"), taken=attempt.get("taken", "") if attempt else "", miss_count=0))
    for row in rows: row["miss_count"] = count_docs("attempt_answers", {"question_id": row["id"], "is_correct": 0, "attempt_id": {"$in": attempt_ids}})
    return render_template("revision.html", rows=rows)

@app.route("/leaderboard")
@login_required
def leaderboard():
    stats = {}
    for attempt in find_many("attempts"):
        user = find_one("users", {"id": attempt["user_id"]})
        if not user or user.get("is_admin"): continue
        stat = stats.setdefault(user["id"], {"name": user["name"], "pts": 0, "tests": 0})
        stat["pts"] += attempt["score"]; stat["tests"] += 1
    rows = sorted([dict(v, pts=round(v["pts"], 2)) for v in stats.values()], key=lambda x: x["pts"], reverse=True)[:20]
    return render_template("leaderboard.html", rows=rows)

# ---------- Live updates ----------
_cache = {"t": 0, "items": []}
def fetch_feed():
    url = app.config["GATE_FEED_URL"]
    if not url: return []
    if time.time() - _cache["t"] < 1800: return _cache["items"]
    try:
        root = ET.fromstring(urllib.request.urlopen(urllib.request.Request(url, headers={"User-Agent": "gate-portal"}), timeout=6).read())
        # Handle RSS items and Atom entries without assuming a particular namespace.
        def child_text(node, name):
            return next(((c.text or "").strip() for c in node if c.tag.rsplit("}", 1)[-1] == name), "")
        items = []
        for node in root.iter():
            if node.tag.rsplit("}", 1)[-1] not in ("item", "entry"):
                continue
            link = child_text(node, "link")
            if not link:
                link_node = next((c for c in node if c.tag.rsplit("}", 1)[-1] == "link"), None)
                link = link_node.attrib.get("href", "") if link_node is not None else ""
            items.append({"title": child_text(node, "title"), "link": link,
                          "date": (child_text(node, "pubDate") or child_text(node, "published") or child_text(node, "updated"))[:25]})
            if len(items) == 10:
                break
        _cache.update(t=time.time(), items=items)
    except Exception:
        _cache["t"] = time.time()
    return _cache["items"]

@app.route("/updates")
@login_required
def updates():
    return render_template("updates.html", items=find_many("updates", sort=[("id", -1)]), feed=fetch_feed(),
                           exam=app.config["GATE_EXAM_DATE"])

# ---------- Admin ----------
@app.route("/admin")
@admin_required
def admin():
    subs = find_many("subjects", sort=[("id", 1)])
    topics = find_many("topics", sort=[("subject_id", 1), ("id", 1)])
    for topic in topics: topic["sname"] = (find_one("subjects", {"id": topic["subject_id"]}) or {}).get("name", "")
    return render_template("admin.html", subs=subs, topics=topics, updates=find_many("updates", sort=[("id", -1)]),
                           nu=count_docs("users"), nq=count_docs("questions"))

@app.post("/admin/<what>")
@admin_required
def admin_add(what):
    f = request.form
    try:
        if what == "update":
            add_doc("updates", {"title": f["title"], "body": f["body"], "link": f["link"], "created": datetime.now().isoformat()})
        elif what == "topic":
            add_doc("topics", {"subject_id": int(f["subject_id"]), "name": f["name"].strip()})
        elif what == "resource":
            if not f["url"].startswith(("http://", "https://")): raise ValueError("URL must start with http(s)")
            add_doc("resources", {"topic_id": int(f["topic_id"]), "title": f["title"], "url": f["url"], "kind": f["kind"]})
        elif what == "question":
            if f["ans"] not in "ABCD": raise ValueError("bad answer")
            subject_id = int(f["subject_id"])
            topic_id = f.get("topic_id") or None
            source_url = f.get("source", "Original practice").strip()[:300]
            if source_url and source_url.startswith(("http://", "https://")) is False and source_url != "Original practice":
                raise ValueError("Source must be a URL or Original practice")
            if topic_id:
                topic_id = int(topic_id)
                topic = find_one("topics", {"id": topic_id})
                if not topic or topic["subject_id"] != subject_id: raise ValueError("Topic must belong to the selected subject")
            add_doc("questions", {"subject_id": subject_id, "topic_id": topic_id, "text": f["text"], "a": f["a"], "b": f["b"], "c": f["c"], "d": f["d"],
                                   "ans": f["ans"], "expl": f["expl"], "source": source_url or "Original practice", "year": int(f["year"]) if f.get("year") else None,
                                   "kind": f.get("kind", "MCQ"), "important": int(bool(f.get("important"))), "bank_id": None, "bank_set": None})
        else: abort(404)
        flash(f"{what.title()} added.", "ok")
    except (KeyError, ValueError) as e:
        flash(f"Could not add: {e}", "err")
    return redirect(url_for("admin"))

@app.post("/admin/import-questions")
@admin_required
def import_questions():
    upload = request.files.get("file")
    if not upload or not upload.filename.lower().endswith(".csv"):
        flash("Choose a .csv question file using the provided template.", "err"); return redirect(url_for("admin"))
    required = {"subject", "topic", "question", "a", "b", "c", "d", "answer", "explanation", "source", "year", "important"}
    added = skipped = 0
    grouped = {}
    try:
        content = upload.stream.read().decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(content))
        if not reader.fieldnames or not required.issubset({h.strip().lower() for h in reader.fieldnames}):
            raise ValueError("CSV headers do not match the downloadable template")
        for line, row in enumerate(reader, start=2):
            row = {str(k).strip().lower(): (v or "").strip() for k, v in row.items() if k}
            subject = next((s for s in find_many("subjects") if s["name"].casefold() == row.get("subject", "").casefold()), None)
            if not subject or not row.get("question") or row.get("answer", "").upper() not in ("A", "B", "C", "D") or any(not row.get(k) for k in ("a", "b", "c", "d")):
                skipped += 1; continue
            topic_id = None
            if row.get("topic"):
                topic = next((t for t in find_many("topics", {"subject_id": subject["id"]}) if t["name"].casefold() == row["topic"].casefold()), None)
                if not topic: skipped += 1; continue
                topic_id = topic["id"]
            year = int(row["year"]) if row.get("year", "").isdigit() else None
            source = row.get("source") or "Original practice"
            if source != "Original practice" and not source.startswith(("https://", "http://")):
                skipped += 1; continue
            grouped.setdefault(subject["id"], []).append((topic_id, row["question"], row["a"], row["b"], row["c"], row["d"], row["answer"].upper(), row.get("explanation", ""), source, year, int(row.get("important", "").lower() in ("1", "yes", "true"))))
        for subject_id, items in grouped.items():
            random.shuffle(items)
            bank_id = uuid.uuid4().hex
            add_doc("question_banks", {"id": bank_id, "subject_id": subject_id, "name": os.path.basename(upload.filename), "created": datetime.now().isoformat(), "total": len(items)})
            for index, item in enumerate(items):
                topic_id, *values = item
                add_doc("questions", {"subject_id": subject_id, "topic_id": topic_id, "text": values[0], "a": values[1], "b": values[2], "c": values[3], "d": values[4],
                                       "ans": values[5], "expl": values[6], "source": values[7], "year": values[8], "kind": "MCQ", "important": values[9],
                                       "bank_id": bank_id, "bank_set": index // 30 + 1})
                added += 1
        flash(f"Imported {added} questions into shuffled sets of up to 30. {skipped} rows were skipped (check subject/topic names, required fields, and answer labels).", "ok" if added else "err")
    except (UnicodeDecodeError, csv.Error, ValueError) as e:
        flash(f"Could not import CSV: {e}", "err")
    return redirect(url_for("admin"))

@app.post("/admin/import-document")
@admin_required
def import_document():
    upload = request.files.get("file")
    if not upload or not upload.filename:
        flash("Choose a PDF or DOCX document.", "err"); return redirect(url_for("admin"))
    data = upload.stream.read(20 * 1024 * 1024 + 1)
    if len(data) > 20 * 1024 * 1024:
        flash("Document is larger than 20 MB. Split it into smaller files and import each one.", "err"); return redirect(url_for("admin"))
    try:
        subject_id = int(request.form.get("subject_id", ""))
        subject = find_one("subjects", {"id": subject_id})
        if not subject: raise ValueError("Choose a valid subject")
        topic_id = request.form.get("topic_id") or None
        if topic_id:
            topic_id = int(topic_id)
            topic = find_one("topics", {"id": topic_id, "subject_id": subject_id})
            if not topic: raise ValueError("The default topic must belong to the selected subject")
        text = read_question_document(upload.filename, data)
        parsed, issues = parse_question_document(text)
        if not parsed: raise ValueError("No importable questions found. " + " ".join(issues[:2]))
        if len(parsed) > 1000: raise ValueError("This document contains over 1,000 detected questions; split it into smaller documents")
        draft_id = uuid.uuid4().hex
        payload = json.dumps({"questions": parsed, "issues": issues, "year": request.form.get("year", "").strip()})
        add_doc("import_drafts", {"id": draft_id, "user_id": g.user["id"], "subject_id": subject_id, "topic_id": topic_id,
                                   "filename": os.path.basename(upload.filename), "payload": payload, "created": datetime.now().isoformat()})
        return redirect(url_for("preview_document", draft_id=draft_id))
    except Exception as e:
        flash(f"Could not read document: {e}", "err")
        return redirect(url_for("admin"))

@app.route("/admin/import-document/<draft_id>")
@admin_required
def preview_document(draft_id):
    draft = find_one("import_drafts", {"id": draft_id, "user_id": g.user["id"]}) or abort(404)
    subject = find_one("subjects", {"id": draft["subject_id"]})
    draft["subject"] = subject["name"] if subject else "Unknown subject"
    payload = json.loads(draft["payload"])
    default_topic = find_one("topics", {"id": draft["topic_id"]}) if draft.get("topic_id") else None
    topics = find_many("topics", {"subject_id": draft["subject_id"]})
    known = {t["name"].casefold() for t in topics}
    for item in payload["questions"]:
        item["resolved_topic"] = item["topic"] if item["topic"].casefold() in known else (default_topic["name"] if default_topic else "")
        if item["topic"] and item["topic"].casefold() not in known:
            payload["issues"].append(f"Question {item['number']}: topic '{item['topic']}' did not match this subject and will use the selected default topic.")
    return render_template("import_preview.html", draft=draft, payload=payload, topics=topics)

@app.post("/admin/import-document/<draft_id>/cancel")
@admin_required
def cancel_document_import(draft_id):
    remove_doc("import_drafts", {"id": draft_id, "user_id": g.user["id"]})
    flash("Document import cancelled; no questions were added.", "ok")
    return redirect(url_for("admin"))

@app.post("/admin/import-document/<draft_id>/confirm")
@admin_required
def confirm_document_import(draft_id):
    draft = find_one("import_drafts", {"id": draft_id, "user_id": g.user["id"]}) or abort(404)
    payload = json.loads(draft["payload"])
    items = payload["questions"]
    for index, item in enumerate(items):
        item["text"] = request.form.get(f"q{index}_text", item["text"]).strip()
        for option in ("a", "b", "c", "d"):
            item[option] = request.form.get(f"q{index}_{option}", item[option]).strip()
        item["ans"] = request.form.get(f"q{index}_ans", item["ans"]).upper()
        item["expl"] = request.form.get(f"q{index}_expl", item["expl"]).strip()
        item["topic"] = request.form.get(f"q{index}_topic", item["topic"]).strip()
        if not item["text"] or any(not item[k] for k in ("a", "b", "c", "d")) or item["ans"] not in ("A", "B", "C", "D"):
            flash(f"Question {item['number']} is missing text, an option, or a valid correct answer. Fix it in preview and confirm again.", "err")
            return redirect(url_for("preview_document", draft_id=draft_id))
    random.shuffle(items)
    bank_id = uuid.uuid4().hex
    bank_name = (request.form.get("name") or draft["filename"]).strip()[:120] or draft["filename"]
    created = datetime.now().isoformat()
    add_doc("question_banks", {"id": bank_id, "subject_id": draft["subject_id"], "name": bank_name, "created": created, "total": len(items)})
    topics = {t["name"].casefold(): t["id"] for t in find_many("topics", {"subject_id": draft["subject_id"]})}
    for index, item in enumerate(items):
        topic_id = topics.get(item["topic"].casefold()) if item["topic"] else draft["topic_id"]
        if item["topic"] and not topic_id:
            topic_id = draft["topic_id"]
        add_doc("questions", {"subject_id": draft["subject_id"], "topic_id": topic_id, "text": item["text"], "a": item["a"], "b": item["b"], "c": item["c"], "d": item["d"],
                               "ans": item["ans"], "expl": item["expl"], "source": f"Uploaded document: {draft['filename']}", "year": int(payload["year"]) if payload["year"].isdigit() else None,
                               "kind": "MCQ", "important": 0, "bank_id": bank_id, "bank_set": index // 30 + 1})
    remove_doc("import_drafts", {"id": draft_id})
    sets = (len(items) + 29) // 30
    flash(f"Imported {len(items)} questions into {sets} set(s) for {draft['subject_id']}. Questions were shuffled and divided into sets of up to 30.", "ok")
    return redirect(url_for("admin"))

@app.post("/admin/update/<int:uid>/delete")
@admin_required
def del_update(uid):
    remove_doc("updates", {"id": uid}); return redirect(url_for("admin"))

init_db()

@app.get("/health")
def health():
    try:
        mongo_client.admin.command("ping")
        counts = {name: count_docs(name) for name in ("users", "subjects", "topics", "questions", "attempts")}
        return jsonify(status="ok", database=app.config["MONGO_DATABASE"], collections=counts), 200
    except PyMongoError:
        app.logger.exception("MongoDB health check failed")
        return jsonify(status="error", message="MongoDB is unavailable"), 503

if __name__ == "__main__":
    app.run(debug=True)
